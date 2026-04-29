#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文档文本提取工具
支持PDF、DOC、DOCX格式文件的文字提取
"""
import tkinter as tk
from tkinter import ttk, filedialog, scrolledtext, messagebox
import os
import sys
import re


class DocumentTextExtractor:
    def __init__(self, root=None):
        # 如果没有提供根窗口，则创建一个隐藏的根窗口
        self.root = root if root else tk.Tk()
        self.root.withdraw()  # 隐藏根窗口
        
    def extract_text(self, file_path):
        """提取文件文本内容"""
        # 检查文件扩展名
        file_ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if file_ext == '.pdf':
                text = self.extract_pdf_text(file_path)
            elif file_ext in ['.doc', '.docx']:
                text = self.extract_doc_text(file_path)
            else:
                raise Exception(f"不支持的文件格式: {file_ext}")
        except Exception as e:
            raise e
        
        return text
    
    def extract_pdf_text(self, file_path):
        """提取PDF文件文本"""
        try:
            # 尝试导入PyPDF2
            import PyPDF2
            with open(file_path, 'rb') as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                
                return text
        except ImportError:
            raise ImportError("缺少PyPDF2库，无法提取PDF文本。请安装: pip install PyPDF2")
        except Exception as e:
            raise e
    
    def extract_doc_text(self, file_path):
        """提取DOC/DOCX文件文本"""
        try:
            # 尝试导入python-docx
            from docx import Document
            if file_path.lower().endswith('.docx'):
                doc = Document(file_path)
                text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                
                return text
            else:
                # 对于.doc文件，需要使用其他库
                return self.extract_doc_old_format(file_path)
        except ImportError:
            raise ImportError("缺少python-docx库，无法提取DOCX文本。请安装: pip install python-docx")
        except Exception as e:
            raise e
    
    def extract_doc_old_format(self, file_path):
        """提取旧版DOC文件文本"""
        try:
            # 尝试使用pywin32（仅Windows）
            import win32com.client
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(file_path)
            text = doc.Content.Text
            doc.Close()
            word.Quit()
            
            return text
        except ImportError:
            raise ImportError("无法处理DOC文件格式。请安装: pip install pywin32 (仅Windows)\n或者安装: pip install python-docx (仅支持.docx)")
        except Exception as e:
            raise e
    
    def process_text(self, text):
        """处理文本，提取摘要和关键词"""
        if not text:
            return "", ""
        
        # 提取摘要
        abstract = self.extract_abstract(text)
        # 提取关键词
        keywords = self.extract_keywords(text)
        
        return abstract, keywords
    
    def extract_abstract(self, text):
        """提取摘要文本"""
        # 查找摘要开始位置
        abstract_start = -1
        
        # 检查不同形式的摘要标识
        patterns = [
            r'摘要[：:]',  # 匹配"摘要："或"摘要:"
            r'ABSTRACT[：:]',  # 匹配"ABSTRACT："或"ABSTRACT:"
            r'Abstract[：:]',  # 匹配"Abstract："或"Abstract:"
        ]
        
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                abstract_start = match.end()
                break
        
        if abstract_start == -1:
            # 如果没有找到摘要标识，尝试通过上下文判断
            lines = text.split('\n')
            for i, line in enumerate(lines):
                if '摘要' in line or 'abstract' in line.lower():
                    abstract_start = text.find(line)
                    abstract_start += len(line)
                    break
        
        if abstract_start != -1:
            # 查找摘要结束位置（关键词开始处或文本结束）
            keywords_pos = -1
            for keyword_pattern in [r'关键词[：:]', r'KEYWORDS[：:]', r'Keywords[：:]']:
                match = re.search(keyword_pattern, text[abstract_start:], re.IGNORECASE)
                if match:
                    keywords_pos = abstract_start + match.start()
                    break
            
            if keywords_pos != -1:
                abstract = text[abstract_start:keywords_pos].strip()
            else:
                abstract = text[abstract_start:].strip()
            
            # 删除摘要中的换行符号和多余空格，确保所有字都在一行
            abstract = re.sub(r'\n+', ' ', abstract)
            abstract = re.sub(r'\s+', ' ', abstract).strip()
            return abstract
        
        return ""
    
    def extract_keywords(self, text):
        """提取关键词文本"""
        # 查找关键词开始位置
        keywords_start = -1
        
        # 检查不同形式的关键词标识
        patterns = [
            r'关键词[：:]',  # 匹配"关键词："或"关键词:"
            r'KEYWORDS[：:]',  # 匹配"KEYWORDS："或"KEYWORDS:"
            r'Keywords[：:]',  # 匹配"Keywords："或"Keywords:"
        ]
        
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                keywords_start = match.end()
                break
        
        if keywords_start == -1:
            # 如果没有找到关键词标识，尝试通过上下文判断
            lines = text.split('\n')
            for i, line in enumerate(lines):
                if '关键词' in line or 'keyword' in line.lower():
                    keywords_start = text.find(line)
                    keywords_start += len(line)
                    break
        
        if keywords_start != -1:
            # 关键词通常到文本结束或遇到其他段落标题
            # 简单处理：提取到文本结束或遇到常见段落标题
            remaining_text = text[keywords_start:]
            end_patterns = [r'摘要[：:]', r'ABSTRACT[：:]', r'Abstract[：:]', r'正文', r'引言', r'1\.', r'一、']
            
            min_pos = len(remaining_text)
            for pattern in end_patterns:
                match = re.search(pattern, remaining_text, re.IGNORECASE)
                if match and match.start() < min_pos:
                    min_pos = match.start()
            
            keywords = remaining_text[:min_pos].strip()
            # 删除关键词中的换行符号和多余空格，确保所有字都在一行
            keywords = re.sub(r'\n+', ' ', keywords)
            keywords = re.sub(r'\s+', ' ', keywords).strip()
            return keywords
        
        return ""


def extract_document_info(file_path):
    """
    从文档中提取文本、摘要和关键词的便捷函数
    :param file_path: 文档文件路径
    :return: tuple (text, abstract, keywords)
    """
    extractor = DocumentTextExtractor()
    text = extractor.extract_text(file_path)
    abstract, keywords = extractor.process_text(text)
    return text, abstract, keywords


def main():
    """主函数 - 用于独立运行GUI程序"""
    root = tk.Tk()
    app = DocumentTextExtractorGUI(root)
    root.mainloop()


class DocumentTextExtractorGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("文档文本提取工具")
        self.root.geometry("1000x700")
        
        # 创建界面
        self.create_widgets()
        
    def create_widgets(self):
        """创建界面控件"""
        # 主框架
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        # 配置网格权重
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        main_frame.columnconfigure(1, weight=1)
        main_frame.rowconfigure(5, weight=1)
        
        # 标题
        title_label = ttk.Label(main_frame, text="文档文本提取工具", font=("Arial", 16, "bold"))
        title_label.grid(row=0, column=0, columnspan=2, pady=(0, 20))
        
        # 文件选择按钮
        self.select_button = ttk.Button(
            main_frame, 
            text="选择文件", 
            command=self.select_file
        )
        self.select_button.grid(row=1, column=0, columnspan=2, pady=(0, 20))
        
        # 文件路径显示
        self.file_path_label = ttk.Label(main_frame, text="未选择文件", foreground="gray")
        self.file_path_label.grid(row=2, column=0, columnspan=2, sticky=tk.W, pady=(0, 10))
        
        # 摘要文本框
        ttk.Label(main_frame, text="摘要:").grid(row=3, column=0, sticky=tk.W, pady=(0, 5))
        self.abstract_area = scrolledtext.ScrolledText(
            main_frame, 
            wrap=tk.WORD, 
            width=60, 
            height=8
        )
        self.abstract_area.grid(row=4, column=0, sticky=(tk.W, tk.E, tk.N, tk.S), padx=(0, 5))
        
        # 关键词文本框
        ttk.Label(main_frame, text="关键词:").grid(row=3, column=1, sticky=tk.W, pady=(0, 5))
        self.keywords_area = scrolledtext.ScrolledText(
            main_frame, 
            wrap=tk.WORD, 
            width=40, 
            height=8
        )
        self.keywords_area.grid(row=4, column=1, sticky=(tk.W, tk.E, tk.N, tk.S), padx=(5, 0))
        
        # 原始文本显示区域
        ttk.Label(main_frame, text="原始文本内容:").grid(row=5, column=0, columnspan=2, sticky=tk.W, pady=(10, 5))
        self.text_area = scrolledtext.ScrolledText(
            main_frame, 
            wrap=tk.WORD, 
            width=100, 
            height=10
        )
        self.text_area.grid(row=6, column=0, columnspan=2, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(0, 10))
        
        # 清空按钮
        self.clear_button = ttk.Button(
            main_frame, 
            text="清空内容", 
            command=self.clear_text
        )
        self.clear_button.grid(row=7, column=0, pady=(0, 10))
        
        # 退出按钮
        self.exit_button = ttk.Button(
            main_frame, 
            text="退出程序", 
            command=self.exit_program
        )
        self.exit_button.grid(row=7, column=1, pady=(0, 10))
        
        # 配置列权重
        main_frame.columnconfigure(0, weight=2)
        main_frame.columnconfigure(1, weight=1)
        
    def select_file(self):
        """选择文件"""
        file_path = filedialog.askopenfilename(
            title="选择文档文件",
            filetypes=[
                ("PDF文件", "*.pdf"),
                ("Word文档", "*.doc"),
                ("Word文档", "*.docx"),
                ("所有支持的文件", "*.pdf;*.doc;*.docx")
            ]
        )
        
        if file_path:
            self.file_path_label.config(text=f"文件: {os.path.basename(file_path)}")
            self.extract_and_process_file(file_path)
    
    def extract_and_process_file(self, file_path):
        """提取并处理文件"""
        try:
            text = self.extract_text(file_path)
            # 显示原始文本
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, text)
            
            # 自动处理文本，提取摘要和关键词
            self.process_text()
        except Exception as e:
            messagebox.showerror("错误", f"处理文件时出错: {str(e)}")
    
    def extract_text(self, file_path):
        """提取文件文本内容"""
        # 检查文件扩展名
        file_ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if file_ext == '.pdf':
                text = self.extract_pdf_text(file_path)
            elif file_ext in ['.doc', '.docx']:
                text = self.extract_doc_text(file_path)
            else:
                messagebox.showerror("错误", f"不支持的文件格式: {file_ext}")
                return ""
        except Exception as e:
            messagebox.showerror("错误", f"提取文本时出错: {str(e)}")
            return ""
        
        return text
    
    def extract_pdf_text(self, file_path):
        """提取PDF文件文本"""
        try:
            # 尝试导入PyPDF2
            import PyPDF2
            with open(file_path, 'rb') as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                
                return text
        except ImportError:
            messagebox.showwarning(
                "警告", 
                "缺少PyPDF2库，无法提取PDF文本。\n"
                "请安装: pip install PyPDF2"
            )
            return ""
        except Exception as e:
            raise e
    
    def extract_doc_text(self, file_path):
        """提取DOC/DOCX文件文本"""
        try:
            # 尝试导入python-docx
            from docx import Document
            if file_path.lower().endswith('.docx'):
                doc = Document(file_path)
                text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                
                return text
            else:
                # 对于.doc文件，需要使用其他库
                return self.extract_doc_old_format(file_path)
        except ImportError:
            messagebox.showwarning(
                "警告", 
                "缺少python-docx库，无法提取DOCX文本。\n"
                "请安装: pip install python-docx"
            )
            return ""
        except Exception as e:
            raise e
    
    def extract_doc_old_format(self, file_path):
        """提取旧版DOC文件文本"""
        try:
            # 尝试使用pywin32（仅Windows）
            import win32com.client
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(file_path)
            text = doc.Content.Text
            doc.Close()
            word.Quit()
            
            return text
        except ImportError:
            messagebox.showwarning(
                "警告", 
                "无法处理DOC文件格式。\n"
                "请安装: pip install pywin32 (仅Windows)\n"
                "或者安装: pip install python-docx (仅支持.docx)"
            )
            return ""
        except Exception as e:
            raise e
    
    def process_text(self):
        """处理文本，提取摘要和关键词"""
        # 获取原始文本
        raw_text = self.text_area.get(1.0, tk.END).strip()
        
        if not raw_text:
            return
        
        # 提取摘要
        abstract = self.extract_abstract(raw_text)
        # 提取关键词
        keywords = self.extract_keywords(raw_text)
        
        # 显示在对应的文本框中
        self.abstract_area.delete(1.0, tk.END)
        self.abstract_area.insert(tk.END, abstract)
        
        self.keywords_area.delete(1.0, tk.END)
        self.keywords_area.insert(tk.END, keywords)
    
    def extract_abstract(self, text):
        """提取摘要文本"""
        # 查找摘要开始位置
        abstract_start = -1
        
        # 检查不同形式的摘要标识
        patterns = [
            r'摘要[：:]',  # 匹配"摘要："或"摘要:"
            r'ABSTRACT[：:]',  # 匹配"ABSTRACT："或"ABSTRACT:"
            r'Abstract[：:]',  # 匹配"Abstract："或"Abstract:"
        ]
        
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                abstract_start = match.end()
                break
        
        if abstract_start == -1:
            # 如果没有找到摘要标识，尝试通过上下文判断
            lines = text.split('\n')
            for i, line in enumerate(lines):
                if '摘要' in line or 'abstract' in line.lower():
                    abstract_start = text.find(line)
                    abstract_start += len(line)
                    break
        
        if abstract_start != -1:
            # 查找摘要结束位置（关键词开始处或文本结束）
            keywords_pos = -1
            for keyword_pattern in [r'关键词[：:]', r'KEYWORDS[：:]', r'Keywords[：:]']:
                match = re.search(keyword_pattern, text[abstract_start:], re.IGNORECASE)
                if match:
                    keywords_pos = abstract_start + match.start()
                    break
            
            if keywords_pos != -1:
                abstract = text[abstract_start:keywords_pos].strip()
            else:
                abstract = text[abstract_start:].strip()
            
            # 删除摘要中的换行符号和多余空格，确保所有字都在一行
            abstract = re.sub(r'\n+', ' ', abstract)
            abstract = re.sub(r'\s+', ' ', abstract).strip()
            return abstract
        
        return ""
    
    def extract_keywords(self, text):
        """提取关键词文本"""
        # 查找关键词开始位置
        keywords_start = -1
        
        # 检查不同形式的关键词标识
        patterns = [
            r'关键词[：:]',  # 匹配"关键词："或"关键词:"
            r'KEYWORDS[：:]',  # 匹配"KEYWORDS："或"KEYWORDS:"
            r'Keywords[：:]',  # 匹配"Keywords："或"Keywords:"
        ]
        
        for pattern in patterns:
            match = re.search(pattern, text, re.IGNORECASE)
            if match:
                keywords_start = match.end()
                break
        
        if keywords_start == -1:
            # 如果没有找到关键词标识，尝试通过上下文判断
            lines = text.split('\n')
            for i, line in enumerate(lines):
                if '关键词' in line or 'keyword' in line.lower():
                    keywords_start = text.find(line)
                    keywords_start += len(line)
                    break
        
        if keywords_start != -1:
            # 关键词通常到文本结束或遇到其他段落标题
            # 简单处理：提取到文本结束或遇到常见段落标题
            remaining_text = text[keywords_start:]
            end_patterns = [r'摘要[：:]', r'ABSTRACT[：:]', r'Abstract[：:]', r'正文', r'引言', r'1\.', r'一、']
            
            min_pos = len(remaining_text)
            for pattern in end_patterns:
                match = re.search(pattern, remaining_text, re.IGNORECASE)
                if match and match.start() < min_pos:
                    min_pos = match.start()
            
            keywords = remaining_text[:min_pos].strip()
            # 删除关键词中的换行符号和多余空格，确保所有字都在一行
            keywords = re.sub(r'\n+', ' ', keywords)
            keywords = re.sub(r'\s+', ' ', keywords).strip()
            return keywords
        
        return ""
    
    def clear_text(self):
        """清空文本区域"""
        self.text_area.delete(1.0, tk.END)
        self.abstract_area.delete(1.0, tk.END)
        self.keywords_area.delete(1.0, tk.END)
        self.file_path_label.config(text="未选择文件")
    
    def exit_program(self):
        """退出程序"""
        self.root.destroy()


if __name__ == "__main__":
    main()